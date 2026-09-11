"""Ingest: line numbers, offsets, the bibliography split, and the binary formats."""

from __future__ import annotations

import zipfile
from collections.abc import Sequence
from pathlib import Path

import docx
import pymupdf
import pytest

from proofpath.document import Document, Locator, PageError
from proofpath.ingest import (
    IngestError,
    _line_text,
    find_bibliography,
    from_docx,
    from_markdown,
    from_pdf,
    from_plain,
    from_text,
    load,
    paragraphs_from_lines,
    split_references,
)

PLAIN = "Alpha one. Alpha two.\n\nBeta one. Beta two.\n"

WRAPPED = "The first sentence is here. The second sentence\nwraps onto line two.\n"

MARKDOWN = (
    "# Title\n"  # 1
    "\n"  # 2
    "A **bold** claim here. It\n"  # 3
    "continues on line four.\n"  # 4
    "\n"  # 5
    "```python\n"  # 6
    "code = 1  # not prose.\n"  # 7
    "```\n"  # 8
    "\n"  # 9
    "- first item\n"  # 10
    "- second item\n"  # 11
)

NUMBERED_BIBLIOGRAPHY = (
    "Body sentence here.\n"  # 1
    "\n"  # 2
    "## References\n"  # 3
    "\n"  # 4
    "[1] Smith, J. (2020). A title. Journal.\n"  # 5
    "[2] Doe, A. (2021). Another title,\n"  # 6
    "    Journal of Things.\n"  # 7
)


def _check_invariants(document: Document) -> None:
    """Every invariant Task 5.1 asks the builder to guarantee."""
    for position, paragraph in enumerate(document.paragraphs):
        assert paragraph.index == position
        assert paragraph.lines  # never empty
        assert paragraph.lines[0][1] == 0  # the first line starts at offset 0
        assert list(paragraph.lines) == sorted(paragraph.lines, key=lambda item: item[1])
        for sentence in paragraph.sentences:
            assert sentence.text == paragraph.text[sentence.start : sentence.end]
            # `Sentence.locator` uses the same rule as `Document.locate`.
            assert sentence.locator == document.locate(position, sentence.start)


def test_a_blank_line_separates_paragraphs_and_line_numbers_survive() -> None:
    document = Document(
        name="notes.txt",
        kind="text",
        paragraphs=tuple(paragraphs_from_lines(PLAIN.splitlines())),
        references=(),
        pages=1,
    )
    assert [p.text for p in document.paragraphs] == ["Alpha one. Alpha two.", "Beta one. Beta two."]
    assert [p.locator.line for p in document.paragraphs] == [1, 3]
    assert [s.text for s in document.paragraphs[0].sentences] == ["Alpha one.", "Alpha two."]
    assert document.paragraphs[0].sentences[1].locator == Locator(line=1, column=12)
    assert document.paragraphs[1].sentences[0].locator == Locator(line=3, column=1)
    _check_invariants(document)


def test_the_second_sentence_of_a_wrapped_paragraph_is_on_the_second_line() -> None:
    paragraph = paragraphs_from_lines(WRAPPED.splitlines())[0]
    assert paragraph.text == "The first sentence is here. The second sentence wraps onto line two."
    assert paragraph.lines == (
        (1, 0),
        (2, len("The first sentence is here. The second sentence") + 1),
    )
    assert paragraph.sentences[1].locator == Locator(line=1, column=29)
    offset = paragraph.text.index("wraps")
    assert paragraph.lines[1][1] == offset


def test_paragraphs_can_start_at_a_line_offset_and_carry_a_page() -> None:
    paragraph = paragraphs_from_lines(["Only line."], page=4, first_line=112)[0]
    assert paragraph.locator == Locator(line=112, column=1, page=4)
    assert paragraph.sentences[0].locator.page == 4


def test_markdown_strips_prefixes_and_emphasis_without_losing_line_numbers() -> None:
    paragraphs = paragraphs_from_lines(MARKDOWN.splitlines(), markdown=True)
    assert [p.text for p in paragraphs] == [
        "Title",
        "A bold claim here. It continues on line four.",
        "first item",
        "second item",
    ]
    assert [p.locator.line for p in paragraphs] == [1, 3, 10, 11]
    body = paragraphs[1]
    assert body.lines == ((3, 0), (4, len("A bold claim here. It") + 1))
    # The stripped `**` shifts offsets, and line 4 still maps to line 4.
    assert body.sentences[1].locator == Locator(line=3, column=20)
    document = Document(
        name="d.md", kind="markdown", paragraphs=tuple(paragraphs), references=(), pages=1
    )
    assert document.locate(1, body.text.index("continues")) == Locator(line=4, column=1)
    _check_invariants(document)


def test_a_fenced_code_block_is_skipped_entirely() -> None:
    paragraphs = paragraphs_from_lines(MARKDOWN.splitlines(), markdown=True)
    assert all("code = 1" not in p.text for p in paragraphs)


def test_markdown_prefixes_are_only_stripped_in_markdown_mode() -> None:
    paragraphs = paragraphs_from_lines(["# Title", "- one", "- two"])
    assert [p.text for p in paragraphs] == ["# Title - one - two"]


def test_a_blockquote_marker_is_stripped() -> None:
    paragraphs = paragraphs_from_lines(["> Quoted claim here."], markdown=True)
    assert paragraphs[0].text == "Quoted claim here."


def test_a_horizontal_rule_separates_paragraphs_without_becoming_one() -> None:
    paragraphs = paragraphs_from_lines(["Before it.", "---", "After it."], markdown=True)
    assert [p.text for p in paragraphs] == ["Before it.", "After it."]
    assert [p.locator.line for p in paragraphs] == [1, 3]


def test_a_table_row_is_its_own_paragraph() -> None:
    lines = ["Intro line.", "| a | b |", "| - | - |", "| 1 | 2 |"]
    paragraphs = paragraphs_from_lines(lines, markdown=True)
    assert [p.locator.line for p in paragraphs] == [1, 2, 4]
    assert paragraphs[1].text == "a   b"


def test_a_hyphen_at_the_end_of_a_line_is_joined_without_a_space() -> None:
    paragraph = paragraphs_from_lines(["An inter-", "national study ran."])[0]
    assert paragraph.text == "An international study ran."
    assert paragraph.lines == ((1, 0), (2, len("An inter")))
    assert paragraph.text[paragraph.lines[1][1] :] == "national study ran."


def test_a_hyphen_before_an_upper_case_line_keeps_the_hyphen() -> None:
    paragraph = paragraphs_from_lines(["A German-", "American study."])[0]
    assert paragraph.text == "A German-American study."
    assert paragraph.lines == ((1, 0), (2, len("A German-")))


def test_an_empty_input_has_no_paragraphs() -> None:
    assert paragraphs_from_lines([]) == []
    assert paragraphs_from_lines(["   ", "\t"]) == []


def test_an_empty_document_holds_nothing_and_reports_nothing() -> None:
    # An empty file is not a failure: there is nothing to read and nothing to report.
    document = from_text("", name="empty.txt", kind="text")
    assert document.paragraphs == ()
    assert document.references == ()
    assert document.errors == ()
    assert document.pages == 1


def test_find_bibliography_matches_a_heading_case_insensitively() -> None:
    for heading in ("References", "BIBLIOGRAPHY", "Works Cited", "references:", "Reference List"):
        paragraphs = paragraphs_from_lines([heading, "", "[1] An entry."])
        assert find_bibliography(paragraphs) == 0, heading


def test_find_bibliography_accepts_a_section_number() -> None:
    for heading in ("7 References", "7. References", "7.2 References"):
        paragraphs = paragraphs_from_lines([heading, "", "[1] An entry."])
        assert find_bibliography(paragraphs) == 0, heading


def test_find_bibliography_takes_the_last_heading_that_has_entries_after_it() -> None:
    lines = ["References", "", "Running head noise.", "", "References", "", "[1] Real entry."]
    paragraphs = paragraphs_from_lines(lines)
    assert find_bibliography(paragraphs) == 2


def test_find_bibliography_ignores_a_trailing_heading_with_nothing_after_it() -> None:
    paragraphs = paragraphs_from_lines(["Body text.", "", "References"])
    assert find_bibliography(paragraphs) is None


def test_find_bibliography_ignores_a_sentence_that_merely_mentions_references() -> None:
    paragraphs = paragraphs_from_lines(["See the references below.", "", "[1] An entry."])
    assert find_bibliography(paragraphs) is None


def test_bracketed_entries_keep_their_printed_numbers_and_wrapped_lines() -> None:
    document = from_text(NUMBERED_BIBLIOGRAPHY, name="paper.md", kind="markdown", markdown=True)
    assert [p.text for p in document.paragraphs] == ["Body sentence here."]
    assert [r.number for r in document.references] == [1, 2]
    assert document.references[0].raw == "[1] Smith, J. (2020). A title. Journal."
    assert document.references[1].raw == "[2] Doe, A. (2021). Another title, Journal of Things."
    assert [r.locator.line for r in document.references] == [5, 6]
    _check_invariants(document)


def test_numbered_entries_in_plain_text_keep_their_printed_numbers() -> None:
    text = "Body.\n\nReferences\n\n1. Smith, J. A title.\n2) Doe, A. Another title.\n"
    document = from_text(text, name="paper.txt", kind="text")
    assert [r.number for r in document.references] == [1, 2]
    assert [r.raw for r in document.references] == [
        "1. Smith, J. A title.",
        "2) Doe, A. Another title.",
    ]


def test_a_wide_gap_after_the_number_also_starts_an_entry() -> None:
    paragraphs = paragraphs_from_lines(["11  Smith, J. A title.", "12  Doe, A. Another."])
    references = split_references(paragraphs, start=0)
    assert [r.number for r in references] == [11, 12]


def test_unnumbered_entries_fall_back_to_one_per_paragraph() -> None:
    text = "Body.\n\nReferences\n\nSmith, J. A title. Journal.\n\nDoe, A. Another title. Journal.\n"
    document = from_text(text, name="paper.txt", kind="text")
    assert [r.number for r in document.references] == [1, 2]
    assert [r.raw for r in document.references] == [
        "Smith, J. A title. Journal.",
        "Doe, A. Another title. Journal.",
    ]
    assert [r.locator.line for r in document.references] == [5, 7]


def test_a_markdown_numbered_list_bibliography_falls_back_to_ordinals() -> None:
    text = "Body.\n\n## References\n\n1. Smith, J. A title.\n2. Doe, A. Another title.\n"
    document = from_text(text, name="paper.md", kind="markdown", markdown=True)
    assert [r.number for r in document.references] == [1, 2]
    assert [r.raw for r in document.references] == ["Smith, J. A title.", "Doe, A. Another title."]


def test_a_document_without_a_bibliography_keeps_every_paragraph_as_body() -> None:
    document = from_text(PLAIN, name="notes.txt", kind="text")
    assert document.references == ()
    assert len(document.paragraphs) == 2
    assert document.pages == 1
    assert document.kind == "text"
    _check_invariants(document)


def test_a_bibliography_only_document_has_no_body_and_keeps_its_entries() -> None:
    # A reference list pasted on its own: every paragraph belongs to the bibliography,
    # so the body is empty. An empty body is reported as such, never as a clean run.
    text = (
        "References\n"
        "\n"
        "[1] Smith, J. (2020). A title. Journal.\n"
        "[2] Doe, A. (2021). Another title. Journal of Things.\n"
    )
    document = from_text(text, name="refs.txt", kind="text")
    assert document.paragraphs == ()
    assert [reference.number for reference in document.references] == [1, 2]
    assert [reference.raw for reference in document.references] == [
        "[1] Smith, J. (2020). A title. Journal.",
        "[2] Doe, A. (2021). Another title. Journal of Things.",
    ]


def test_from_markdown_and_from_plain_read_files(tmp_path: Path) -> None:
    markdown_path = tmp_path / "paper.md"
    markdown_path.write_text(NUMBERED_BIBLIOGRAPHY, encoding="utf-8")
    document = from_markdown(markdown_path)
    assert document.name == "paper.md"
    assert document.kind == "markdown"
    assert len(document.references) == 2

    plain_path = tmp_path / "notes.txt"
    plain_path.write_text(PLAIN, encoding="utf-8")
    plain = from_plain(plain_path)
    assert plain.name == "notes.txt"
    assert plain.kind == "text"
    assert len(plain.paragraphs) == 2


def test_undecodable_bytes_are_replaced_rather_than_raised(tmp_path: Path) -> None:
    path = tmp_path / "broken.txt"
    path.write_bytes(b"Caf\xe9 study ran.\n")
    assert "study ran." in from_plain(path).paragraphs[0].text


def test_a_missing_file_raises_ingest_error(tmp_path: Path) -> None:
    with pytest.raises(IngestError):
        from_plain(tmp_path / "absent.txt")


def test_a_bibliography_heading_needs_no_blank_line_under_it() -> None:
    text = "Body sentence.\nReferences\n[1] Entry one.\n[2] Entry two.\n"
    document = from_text(text, name="paper.txt", kind="text")
    assert [p.text for p in document.paragraphs] == ["Body sentence."]
    assert [r.number for r in document.references] == [1, 2]
    assert [r.raw for r in document.references] == ["[1] Entry one.", "[2] Entry two."]
    assert [r.locator.line for r in document.references] == [3, 4]
    _check_invariants(document)


def test_a_markdown_bibliography_heading_needs_no_blank_line_under_it() -> None:
    text = "Body sentence.\n## References\n[1] Entry one.\n[2] Entry two.\n"
    document = from_text(text, name="paper.md", kind="markdown", markdown=True)
    assert [p.text for p in document.paragraphs] == ["Body sentence."]
    assert [r.number for r in document.references] == [1, 2]
    assert [r.raw for r in document.references] == ["[1] Entry one.", "[2] Entry two."]


def test_a_bibliography_heading_is_a_paragraph_of_its_own_in_every_mode() -> None:
    lines = ["Body sentence.", "References", "[1] Entry one."]
    for markdown in (False, True):
        paragraphs = paragraphs_from_lines(lines, markdown=markdown)
        assert [p.text for p in paragraphs] == lines, markdown
        assert [p.locator.line for p in paragraphs] == [1, 2, 3], markdown


def test_text_before_the_first_numbered_entry_is_dropped() -> None:
    text = "Body.\n\nReferences\nSee also the appendix.\n[1] A.\n[2] B.\n"
    document = from_text(text, name="paper.txt", kind="text")
    assert [p.text for p in document.paragraphs] == ["Body."]
    assert [r.number for r in document.references] == [1, 2]
    assert [r.raw for r in document.references] == ["[1] A.", "[2] B."]
    assert [r.locator.line for r in document.references] == [5, 6]


def test_a_line_after_an_entry_still_continues_it() -> None:
    # Only the run *before* the first marker is residue; everything after belongs.
    paragraphs = paragraphs_from_lines(["Notes", "[1] A title,", "a journal, 2020.", "[2] B."])
    references = split_references(paragraphs, start=0)
    assert [r.raw for r in references] == ["[1] A title, a journal, 2020.", "[2] B."]


def test_an_unclosed_code_fence_does_not_swallow_the_rest_of_the_document() -> None:
    text = (
        "Intro line.\n"  # 1
        "\n"  # 2
        "```python\n"  # 3 - never closed
        "code = 1\n"  # 4
        "\n"  # 5
        "A body claim here.\n"  # 6
        "\n"  # 7
        "## References\n"  # 8
        "\n"  # 9
        "[1] Smith, J. A title.\n"  # 10
    )
    document = from_text(text, name="paper.md", kind="markdown", markdown=True)
    body = [p.text for p in document.paragraphs]
    assert "Intro line." in body
    assert "A body claim here." in body
    assert [r.raw for r in document.references] == ["[1] Smith, J. A title."]
    # Nothing is dropped: the lines held inside the unmatched fence come back as text.
    assert any("code = 1" in text for text in body)
    # ... but the fence line itself is markup, not content.
    assert not any("python" in text for text in body)
    _check_invariants(document)


def test_a_closed_code_fence_is_still_skipped_entirely() -> None:
    text = "Intro.\n\n```\ncode = 1\n```\n\nAfter it.\n"
    document = from_text(text, name="paper.md", kind="markdown", markdown=True)
    assert [p.text for p in document.paragraphs] == ["Intro.", "After it."]


def test_a_continuation_line_starting_with_a_year_is_not_a_new_entry() -> None:
    text = (
        "Body.\n"  # 1
        "\n"  # 2
        "References\n"  # 3
        "\n"  # 4
        "[1] Smith, J. A title that wraps,\n"  # 5
        "2020. Journal of Things.\n"  # 6
        "[2] Doe, A. Another title.\n"  # 7
    )
    document = from_text(text, name="paper.txt", kind="text")
    assert [r.number for r in document.references] == [1, 2]
    assert document.references[0].raw == (
        "[1] Smith, J. A title that wraps, 2020. Journal of Things."
    )
    assert document.references[1].raw == "[2] Doe, A. Another title."


def test_an_out_of_sequence_number_continues_the_entry_above_it() -> None:
    paragraphs = paragraphs_from_lines(["1. Smith, J. A title,", "2020. Journal.", "2. Doe, A. B."])
    references = split_references(paragraphs, start=0)
    assert [r.number for r in references] == [1, 2]
    assert references[0].raw == "1. Smith, J. A title, 2020. Journal."


def test_a_number_lost_to_ocr_still_starts_an_entry() -> None:
    paragraphs = paragraphs_from_lines(["[1] A title.", "[3] Another title.", "[4] A third."])
    assert [r.number for r in split_references(paragraphs, start=0)] == [1, 3, 4]


def test_bracketed_entries_may_skip_any_number_of_places() -> None:
    # "[5]" can be nothing but a marker, so a gap in the numbering is not our business.
    paragraphs = paragraphs_from_lines(["[1] A.", "[2] B.", "[5] E."])
    references = split_references(paragraphs, start=0)
    assert [r.number for r in references] == [1, 2, 5]
    assert [r.raw for r in references] == ["[1] A.", "[2] B.", "[5] E."]


def test_bare_entries_may_skip_a_few_places() -> None:
    paragraphs = paragraphs_from_lines(["1. A.", "2. B.", "5. E."])
    references = split_references(paragraphs, start=0)
    assert [r.number for r in references] == [1, 2, 5]
    assert [r.raw for r in references] == ["1. A.", "2. B.", "5. E."]


def test_a_dangling_hyphen_keeps_its_spacing() -> None:
    # Only a hyphen attached to a word breaks a word across lines.
    for continuation in ("Results follow.", "results follow."):
        paragraph = paragraphs_from_lines(["The trial -", continuation])[0]
        assert paragraph.text == f"The trial - {continuation}", continuation


# --- PDF and docx -------------------------------------------------------------
# Binary fixtures are built here, at test time, and never committed.

Blocks = Sequence[tuple[tuple[float, float], str]]

HEADER: tuple[tuple[float, float], str] = ((72, 60), "Preprint under review")
FOOTER: tuple[tuple[float, float], str] = ((72, 760), "Journal of Things")


def _write_pdf(path: Path, pages: Sequence[Blocks]) -> None:
    """Write a PDF whose pages hold text blocks at the given points."""
    document = pymupdf.open()
    try:
        for blocks in pages:
            page = document.new_page()
            for point, text in blocks:
                page.insert_text(point, text)
        document.save(str(path))
    finally:
        document.close()


# A string is one Word paragraph; a list of rows is one table.
DocxItem = str | Sequence[Sequence[str]]


def _write_docx(path: Path, items: Sequence[DocxItem]) -> None:
    """Write a .docx from paragraphs and tables, in the order given."""
    document = docx.Document()
    for item in items:
        if isinstance(item, str):
            document.add_paragraph(item)
            continue
        table = document.add_table(rows=len(item), cols=len(item[0]))
        for row, cells in zip(table.rows, item, strict=True):
            for cell, text in zip(row.cells, cells, strict=True):
                cell.text = text
    document.save(str(path))


def _paper_pdf(path: Path) -> None:
    """Three pages: body, a bibliography that runs on, and furniture on every page."""
    _write_pdf(
        path,
        [
            [
                HEADER,
                ((72, 100), "Alpha one. Alpha two.\nAlpha three."),
                FOOTER,
                ((300, 780), "1"),
            ],
            [
                HEADER,
                ((72, 100), "Beta one is here."),
                ((72, 200), "References"),
                (
                    (72, 260),
                    "[1] Smith, J. (2020). A title. Journal.\n[2] Doe, A. (2021). Another.",
                ),
                FOOTER,
                ((300, 780), "2"),
            ],
            [
                HEADER,
                ((72, 100), "[3] Roe, B. (2022). A third title."),
                FOOTER,
                ((300, 780), "3"),
            ],
        ],
    )


def test_pdf_paragraphs_carry_their_page_and_line_numbers_restart_per_page(
    tmp_path: Path,
) -> None:
    path = tmp_path / "paper.pdf"
    _paper_pdf(path)
    document = from_pdf(path)
    assert document.name == "paper.pdf"
    assert document.kind == "pdf"
    assert document.pages == 3
    assert document.errors == ()
    assert [p.text for p in document.paragraphs] == [
        "Alpha one. Alpha two. Alpha three.",
        "Beta one is here.",
    ]
    # Line 2 of page 2, not line 5 of the file: a PDF locator counts within its page.
    # It counts every line printed on it, too: the running head is dropped as
    # furniture but it was still line 1, so the body starts at line 2.
    assert [p.locator for p in document.paragraphs] == [
        Locator(line=2, column=1, page=1),
        Locator(line=2, column=1, page=2),
    ]
    assert document.paragraphs[0].lines == ((2, 0), (3, len("Alpha one. Alpha two.") + 1))
    _check_invariants(document)


def test_a_footer_repeated_on_three_pages_is_dropped_everywhere(tmp_path: Path) -> None:
    path = tmp_path / "paper.pdf"
    _paper_pdf(path)
    document = from_pdf(path)
    texts = [p.text for p in document.paragraphs] + [r.raw for r in document.references]
    assert all("Journal of Things" not in text for text in texts)
    assert all("Preprint under review" not in text for text in texts)
    # The bare page numbers go too, and they never become a reference.
    assert all(text not in {"1", "2", "3"} for text in texts)


def test_a_line_repeated_on_only_two_pages_is_still_content(tmp_path: Path) -> None:
    path = tmp_path / "short.pdf"
    _write_pdf(
        path,
        [
            [((72, 100), "Alpha one."), FOOTER],
            [((72, 100), "Beta one."), FOOTER],
        ],
    )
    document = from_pdf(path)
    assert [p.text for p in document.paragraphs].count("Journal of Things") == 2


def test_a_long_block_repeated_on_every_page_is_not_furniture(tmp_path: Path) -> None:
    long_line = "This sentence is long enough to be prose and not a running head at all."
    path = tmp_path / "long.pdf"
    _write_pdf(path, [[((72, 700), long_line)] for _ in range(3)])
    document = from_pdf(path)
    assert [p.text for p in document.paragraphs] == [long_line] * 3


def test_a_digits_only_block_in_the_body_band_is_kept(tmp_path: Path) -> None:
    # "42" halfway down the page is a result, not a page number.
    path = tmp_path / "number.pdf"
    _write_pdf(path, [[((72, 400), "42")]])
    document = from_pdf(path)
    assert [p.text for p in document.paragraphs] == ["42"]


def test_a_page_number_in_the_margin_is_dropped_without_repeating(tmp_path: Path) -> None:
    path = tmp_path / "number.pdf"
    _write_pdf(path, [[((72, 100), "A body sentence."), ((300, 780), "42")]])
    document = from_pdf(path)
    assert [p.text for p in document.paragraphs] == ["A body sentence."]


def test_pdf_references_keep_their_printed_numbers_and_page_locators(tmp_path: Path) -> None:
    path = tmp_path / "paper.pdf"
    _paper_pdf(path)
    document = from_pdf(path)
    assert [r.number for r in document.references] == [1, 2, 3]
    assert document.references[0].raw == "[1] Smith, J. (2020). A title. Journal."
    assert [r.locator for r in document.references] == [
        Locator(line=4, column=1, page=2),
        Locator(line=5, column=1, page=2),
        Locator(line=2, column=1, page=3),
    ]


def test_a_bibliography_entry_split_by_a_page_break_becomes_one_reference(
    tmp_path: Path,
) -> None:
    path = tmp_path / "split.pdf"
    _write_pdf(
        path,
        [
            [((72, 100), "References"), ((72, 160), "[1] Smith, J. A title that wraps,")],
            [
                ((72, 100), "across the page break. Journal."),
                ((72, 160), "[2] Doe, A. Another title."),
            ],
        ],
    )
    document = from_pdf(path)
    assert [r.raw for r in document.references] == [
        "[1] Smith, J. A title that wraps, across the page break. Journal.",
        "[2] Doe, A. Another title.",
    ]
    assert document.references[0].locator == Locator(line=2, column=1, page=1)


def test_a_word_broken_by_a_page_break_is_rejoined_in_the_entry(tmp_path: Path) -> None:
    path = tmp_path / "hyphen.pdf"
    _write_pdf(
        path,
        [
            [((72, 100), "References"), ((72, 160), "[1] Smith, J. An inter-")],
            [((72, 100), "national study. Journal.")],
        ],
    )
    document = from_pdf(path)
    # Rejoining the two halves is what dehyphenates them; pasting them would leave
    # "inter- national", which no resolver would match.
    assert [r.raw for r in document.references] == [
        "[1] Smith, J. An international study. Journal."
    ]


def test_a_body_paragraph_never_merges_across_a_page_break(tmp_path: Path) -> None:
    path = tmp_path / "body.pdf"
    _write_pdf(
        path,
        [
            [((72, 100), "A body sentence that runs on,")],
            [((72, 100), "and ends on the next page.")],
        ],
    )
    document = from_pdf(path)
    assert [p.text for p in document.paragraphs] == [
        "A body sentence that runs on,",
        "and ends on the next page.",
    ]
    assert [p.locator.page for p in document.paragraphs] == [1, 2]
    _check_invariants(document)


def test_a_pdf_that_cannot_be_opened_at_all_raises_ingest_error(tmp_path: Path) -> None:
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"%PDF-1.4 garbage")
    with pytest.raises(IngestError, match=r"broken\.pdf"):
        from_pdf(path)
    with pytest.raises(IngestError, match=r"absent\.pdf"):
        from_pdf(tmp_path / "absent.pdf")


def _write_image_pdf(path: Path, *, blank_page: bool = False) -> None:
    """Page 1 holds text, page 2 holds nothing but an image, page 3 (optional) is blank."""
    document = pymupdf.open()
    try:
        document.new_page().insert_text((72, 100), "Alpha one.")
        pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 4, 4), 0)
        document.new_page().insert_image(pymupdf.Rect(72, 72, 172, 172), pixmap=pixmap)
        if blank_page:
            document.new_page()
        document.save(str(path))
    finally:
        document.close()


def test_a_page_that_is_an_image_and_no_text_is_reported_as_such(tmp_path: Path) -> None:
    # A scanned page holds a claim a reader can see and this tool cannot: saying so is
    # the difference between a page with no text and a page nobody could read
    # (product rule 2). OCR is not in v0.1, so the page is named, not guessed at.
    path = tmp_path / "scanned.pdf"
    _write_image_pdf(path)
    document = from_pdf(path)
    assert document.pages == 2
    assert [p.text for p in document.paragraphs] == ["Alpha one."]
    assert document.errors == (
        PageError(page=2, detail="no extractable text; the page is an image (scanned?)"),
    )


def test_a_blank_page_is_not_reported_as_an_image(tmp_path: Path) -> None:
    # A separator page carries neither text nor image: there is nothing to report.
    path = tmp_path / "scanned.pdf"
    _write_image_pdf(path, blank_page=True)
    document = from_pdf(path)
    assert document.pages == 3
    assert [error.page for error in document.errors] == [2]


def test_an_unparseable_page_is_reported_and_the_other_pages_survive(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "paper.pdf"
    _write_pdf(
        path,
        [
            [((72, 100), "Alpha one.")],
            [((72, 100), "Beta one.")],
            [((72, 100), "Gamma one.")],
        ],
    )
    extract = pymupdf.Page.get_text

    def explode(self: pymupdf.Page, *args: object, **kwargs: object) -> object:
        if self.number == 1:  # `Page.number` is 0-based: this is page 2
            raise RuntimeError("broken page")
        return extract(self, *args, **kwargs)

    monkeypatch.setattr(pymupdf.Page, "get_text", explode)
    document = from_pdf(path)
    assert [p.text for p in document.paragraphs] == ["Alpha one.", "Gamma one."]
    assert document.errors == (PageError(page=2, detail="RuntimeError: broken page"),)
    # The page is still counted: the report has to be able to say what it missed.
    assert document.pages == 3


def test_a_pdf_whose_every_page_fails_reports_every_page(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "paper.pdf"
    _write_pdf(path, [[((72, 100), "Alpha one.")], [((72, 100), "Beta one.")]])

    def explode(self: pymupdf.Page, *args: object, **kwargs: object) -> object:
        raise RuntimeError("broken page")

    monkeypatch.setattr(pymupdf.Page, "get_text", explode)
    document = from_pdf(path)
    # An empty document is not a clean one: the coverage is on the record.
    assert document.paragraphs == ()
    assert document.references == ()
    assert [error.page for error in document.errors] == [1, 2]
    assert document.pages == 2


def test_from_docx_splits_on_blank_paragraphs_and_extracts_references(tmp_path: Path) -> None:
    path = tmp_path / "paper.docx"
    _write_docx(
        path,
        [
            "First body paragraph.",
            "",
            "Second body paragraph.",
            "References",
            "[1] Smith, J. (2020). A title. Journal.",
        ],
    )
    document = from_docx(path)
    assert document.name == "paper.docx"
    assert document.kind == "docx"
    assert document.pages == 1
    assert [p.text for p in document.paragraphs] == [
        "First body paragraph.",
        "Second body paragraph.",
    ]
    # An empty paragraph counts as a line, so the numbers are the ones a reader sees.
    assert [p.locator for p in document.paragraphs] == [Locator(line=1), Locator(line=3)]
    assert [p.locator.page for p in document.paragraphs] == [None, None]
    assert [r.raw for r in document.references] == ["[1] Smith, J. (2020). A title. Journal."]
    assert document.references[0].locator == Locator(line=5)
    _check_invariants(document)


def test_a_docx_that_cannot_be_opened_raises_ingest_error(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a zip archive")
    with pytest.raises(IngestError, match=r"broken\.docx"):
        from_docx(path)
    with pytest.raises(IngestError, match=r"absent\.docx"):
        from_docx(tmp_path / "absent.docx")


def test_consecutive_word_paragraphs_stay_separate_paragraphs(tmp_path: Path) -> None:
    # Word puts no blank paragraph between paragraphs; the spacing is a style. Each
    # one is its own block, or a whole page of prose would flatten into one claim.
    path = tmp_path / "paper.docx"
    _write_docx(path, ["First one here.", "Second one here.", "Third one here."])
    document = from_docx(path)
    assert [p.text for p in document.paragraphs] == [
        "First one here.",
        "Second one here.",
        "Third one here.",
    ]
    assert [p.locator for p in document.paragraphs] == [
        Locator(line=1),
        Locator(line=2),
        Locator(line=3),
    ]
    _check_invariants(document)


def test_a_docx_table_becomes_one_paragraph_per_row(tmp_path: Path) -> None:
    path = tmp_path / "table.docx"
    _write_docx(
        path,
        [
            "Before the table.",
            [["Drug", "Effect"], ["Aspirin", "Lowers risk"]],
            "After the table.",
        ],
    )
    document = from_docx(path)
    assert [p.text for p in document.paragraphs] == [
        "Before the table.",
        "Drug | Effect",
        "Aspirin | Lowers risk",
        "After the table.",
    ]
    # The rows keep their place in the one ordinal sequence.
    assert [p.locator.line for p in document.paragraphs] == [1, 2, 3, 4]
    _check_invariants(document)


def test_a_zip_that_is_not_a_word_package_raises_ingest_error(tmp_path: Path) -> None:
    path = tmp_path / "fake.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("hello.txt", "not a Word document")
    with pytest.raises(IngestError, match=r"fake\.docx"):
        from_docx(path)


def test_load_dispatches_on_the_suffix(tmp_path: Path) -> None:
    (tmp_path / "notes.txt").write_text(PLAIN, encoding="utf-8")
    (tmp_path / "paper.md").write_text(NUMBERED_BIBLIOGRAPHY, encoding="utf-8")
    (tmp_path / "paper.markdown").write_text(NUMBERED_BIBLIOGRAPHY, encoding="utf-8")
    _write_docx(tmp_path / "paper.docx", ["A body paragraph."])
    _paper_pdf(tmp_path / "paper.PDF")  # the suffix is matched case-insensitively
    assert load(tmp_path / "notes.txt").kind == "text"
    assert load(tmp_path / "paper.md").kind == "markdown"
    assert load(tmp_path / "paper.markdown").kind == "markdown"
    assert load(tmp_path / "paper.docx").kind == "docx"
    assert load(tmp_path / "paper.PDF").kind == "pdf"


def test_load_names_the_suffix_it_cannot_read(tmp_path: Path) -> None:
    path = tmp_path / "data.xyz"
    path.write_text("Some text.\n", encoding="utf-8")
    with pytest.raises(IngestError, match=r"unsupported file type: \.xyz"):
        load(path)


def test_load_rejects_a_missing_file(tmp_path: Path) -> None:
    with pytest.raises(IngestError, match=r"absent\.md"):
        load(tmp_path / "absent.md")


# --- superscript citation markers (product rule 6: a citation must not be invisible) ---


def _span(text: str, *, superscript: bool = False) -> dict[str, object]:
    """One pymupdf span dict, reduced to the two keys ``_line_text`` reads."""
    return {"text": text, "flags": 1 if superscript else 0}


def test_a_superscript_number_is_spelled_out_as_a_bracketed_marker() -> None:
    spans = [_span("structure prediction"), _span("51", superscript=True), _span(". In parallel")]
    assert _line_text(spans) == "structure prediction[51]. In parallel"


def test_a_superscript_list_or_range_becomes_one_marker() -> None:
    assert _line_text([_span("methods"), _span("3,4", superscript=True)]) == "methods[3,4]"
    en_dash = "\u2013"
    assert (
        _line_text([_span("network"), _span(f"4{en_dash}6", superscript=True)])
        == f"network[4{en_dash}6]"
    )


def test_a_superscript_after_a_digit_is_an_exponent_and_is_left_alone() -> None:
    assert _line_text([_span("10"), _span("3", superscript=True)]) == "103"


def test_a_superscript_on_a_name_that_ends_in_digits_is_still_a_citation() -> None:
    # "Uniclust30", "PDB70", "6Y4F": the digits belong to the name, so a superscript after
    # them is a citation. Only a standalone number carries an exponent.
    assert _line_text([_span("Uniclust30"), _span("36", superscript=True)]) == "Uniclust30[36]"
    assert _line_text([_span("6Y4F"), _span("77", superscript=True)]) == "6Y4F[77]"


def test_a_superscript_after_a_version_string_is_a_citation() -> None:
    # "OpenMM v.7.3.1__69__" in the AlphaFold methods: the trailing "1" is the last field
    # of a version, not a base, so the superscript on it is a citation. A standalone
    # number still takes its exponent.
    spans = [_span("OpenMM v.7.3.1"), _span("69", superscript=True)]
    assert _line_text(spans) == "OpenMM v.7.3.1[69]"
    assert _line_text([_span("10"), _span("3", superscript=True)]) == "103"


def test_a_superscript_after_a_decimal_is_read_as_a_citation() -> None:
    # The cost of the rule above: a digit that follows a decimal point no longer takes an
    # exponent, so "0.5__3__" is read as a citation. Losing a real citation on a version
    # string is the worse trade, and a false marker is one a reader can see.
    assert _line_text([_span("0.5"), _span("3", superscript=True)]) == "0.5[3]"


def test_a_superscript_after_a_space_is_a_citation_not_an_exponent() -> None:
    assert _line_text([_span("the method "), _span("12", superscript=True)]) == "the method [12]"


def test_a_number_the_pdf_does_not_mark_as_superscript_is_left_alone() -> None:
    assert _line_text([_span("prediction"), _span("51")]) == "prediction51"


def test_a_superscript_that_is_not_a_number_is_left_alone() -> None:
    assert _line_text([_span("note"), _span("a", superscript=True)]) == "notea"
    assert _line_text([_span("cost"), _span("1a", superscript=True)]) == "cost1a"


def test_a_line_of_plain_spans_is_joined_unchanged() -> None:
    assert _line_text([_span("one "), _span("two "), _span("three")]) == "one two three"


# --- a bare number that is not an entry number -----------------------------------------


def test_an_author_year_bibliography_is_numbered_by_position_not_by_year() -> None:
    text = (
        "Body sentence here.\n"  # 1
        "\n"  # 2
        "References\n"  # 3
        "\n"  # 4
        "Jacob Devlin, Ming-Wei Chang, Kenton Lee, and Kristina Toutanova.\n"  # 5
        "2019. BERT: pre-training of deep bidirectional transformers. In NAACL.\n"  # 6
        "\n"  # 7
        "Zhilin Yang, Zihang Dai, and Yiming Yang.\n"  # 8
        "2019. XLNet: generalized autoregressive pretraining. In NeurIPS.\n"  # 9
    )
    document = from_text(text, name="paper.txt", kind="text")
    assert [r.number for r in document.references] == [1, 2]
    assert document.references[0].raw.startswith("Jacob Devlin")
    assert document.references[1].raw.startswith("Zhilin Yang")


def test_a_four_digit_number_never_starts_a_bare_entry() -> None:
    paragraphs = paragraphs_from_lines(["2019. BERT: a title.", "", "2019. XLNet: another title."])
    references = split_references(paragraphs, start=0)
    assert [r.number for r in references] == [1, 2]


def test_a_bibliography_that_starts_mid_list_keeps_its_printed_numbers() -> None:
    # A references section that begins at 11 -- a split file, a repeated heading, one page
    # of the list on its own -- must keep its printed numbers: renumbering it by position
    # would point "[12]" at the entry printed 11 and check the claim against the wrong
    # source, without a word about it (product rules 1 and 2).
    text = (
        "A body sentence citing the twelfth entry [12].\n"
        "\n"
        "References\n"
        "\n"
        "11  Kim, S. A title. Journal, 2018.\n"
        "12  Lee, J. Another title. Journal, 2019.\n"
        "13  Park, H. A third title. Journal, 2020.\n"
    )
    document = from_text(text, name="paper.txt", kind="text")
    assert [r.number for r in document.references] == [11, 12, 13]
    assert document.references[1].raw.startswith("12 Lee")


def test_a_bare_list_starting_at_three_after_an_ocr_loss_is_still_numbered() -> None:
    paragraphs = paragraphs_from_lines(["3. Smith, J. A title.", "4. Doe, A. Another title."])
    assert [r.number for r in split_references(paragraphs, start=0)] == [3, 4]


def test_a_bracketed_list_may_start_anywhere() -> None:
    # "[12]" can be nothing but a marker, so the rule for bare numbers does not apply.
    paragraphs = paragraphs_from_lines(["[12] Smith, J. A title.", "[13] Doe, A. Another title."])
    assert [r.number for r in split_references(paragraphs, start=0)] == [12, 13]


def test_a_superscript_after_a_full_stop_is_a_footnote_marker_and_is_left_alone() -> None:
    # ACL and Chicago put a footnote marker after the stop ("...and code.1"); Vancouver
    # and Nature bind a citation to the word before it ("...prediction51."). Rewriting
    # the footnote would cite reference 1 for a sentence that never named it.
    spans = [_span("We release our models and code."), _span("1", superscript=True)]
    assert _line_text(spans) == "We release our models and code.1"


def test_a_superscript_after_a_comma_is_a_footnote_marker_too() -> None:
    spans = [_span("easier to parallelize via distributed training,"), _span("8", superscript=True)]
    assert _line_text(spans) == "easier to parallelize via distributed training,8"


def test_a_superscript_after_a_closing_quote_is_still_a_footnote_marker() -> None:
    # The quote closes the sentence; the stop under it is what the superscript follows.
    spans = [_span("she said so herself.\u201d"), _span("1", superscript=True)]
    assert _line_text(spans) == "she said so herself.\u201d1"


def test_a_superscript_after_a_closing_bracket_is_still_a_citation() -> None:
    spans = [_span("the Protein Data Bank (PDB)"), _span("5", superscript=True)]
    assert _line_text(spans) == "the Protein Data Bank (PDB)[5]"
