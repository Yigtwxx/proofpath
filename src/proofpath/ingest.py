"""Ingest: a file turned into paragraphs that stay addressable.

Spec section 9 step 1. The point of this module is that nothing moves silently:
a paragraph is flattened to one string for claim extraction and retrieval, but
every line it came from keeps its number and its offset in that string, so a
verdict can still say "L112". Markdown prefixes and emphasis markers are stripped
from the text and the offsets are recomputed rather than guessed.

The bibliography is split off here and kept as raw strings only; parsing a
reference into fields is ``resolve.py``'s job, not ours.

PDF and .docx go through the same code: a PDF page's text blocks are handed to
``paragraphs_from_lines`` like any other lines, so dehyphenation and sentence spans
have one implementation. A PDF page that cannot be read becomes a ``PageError`` and
the rest of the document is still ingested (spec section 15); a file that cannot be
opened at all is an ``IngestError``.
"""

from __future__ import annotations

import re
import zipfile
from bisect import bisect_right
from collections.abc import Callable, Sequence
from dataclasses import dataclass, replace
from itertools import pairwise
from pathlib import Path
from typing import TYPE_CHECKING, Any, Literal
from urllib.parse import urlsplit

import docx
import pymupdf
from docx.opc.exceptions import OpcError
from docx.table import Table

from proofpath import resolve
from proofpath.document import (
    Document,
    Kind,
    Locator,
    PageError,
    Paragraph,
    Reference,
    Sentence,
)
from proofpath.retrieval import sentence_spans

if TYPE_CHECKING:  # a post is read by a provider, which imports this module back
    from proofpath.providers.social import Post as SocialPost


class IngestError(ValueError):
    """A document that could not be read. Reported, never returned as an empty one."""


# Markdown line shapes. Only what changes the block structure is recognised; the
# rest of the syntax is left to the emphasis stripper below.
_FENCE = re.compile(r"^\s{0,3}(?:```|~~~)")
_RULE = re.compile(r"^\s{0,3}(?:[-*_=]\s*){3,}$")
_HEADING = re.compile(r"^\s{0,3}#{1,6}\s*")
_HEADING_TAIL = re.compile(r"\s*#+\s*$")
_LIST_ITEM = re.compile(r"^\s{0,3}(?:[-*+]\s+|\d{1,3}[.)]\s+)")
_BLOCKQUOTE = re.compile(r"^\s{0,3}(?:>\s?)+")
_TABLE_ROW = re.compile(r"^\s{0,3}\|")
# ``*``, backticks and ``~~`` always go; ``_`` only where it is not inside a word,
# so snake_case identifiers survive.
_EMPHASIS = re.compile(r"\*+|`+|~~|(?<!\w)_+|_+(?!\w)")
# A hyphen that breaks a word across lines is attached to that word.
_WORD_HYPHEN = re.compile(r"\w-$")

# A markdown heading prefix is accepted in every mode: a draft piped through stdin is
# read as plain text and must still find its "## References".
_BIBLIOGRAPHY_HEADING = re.compile(
    r"^(?:#{1,6}\s+)?(?:\d+(?:\.\d+)*\.?\s*)?"
    r"(?:references|bibliography|works cited|literature cited|reference list"
    r"|references and notes)\s*:?$",
    re.IGNORECASE,
)


Block = Literal["text", "heading", "item", "row"]
# A heading or a table row is a block of its own: nothing may be appended to it.
_STANDALONE: tuple[Block, ...] = ("heading", "row")


def _block_of(line: str) -> Block:
    """Classify a markdown line by what it does to the paragraph structure."""
    if _HEADING.match(line):
        return "heading"
    if _TABLE_ROW.match(line):
        return "row"
    if _LIST_ITEM.match(line):
        return "item"
    return "text"


def _clean(line: str, *, block: Block, markdown: bool) -> str:
    """The text of one line: markdown markers gone, surrounding whitespace gone."""
    if not markdown:
        return line.strip()
    if block == "row":
        # Cells become plain text; the pipes carry no meaning once flattened.
        line = line.strip().strip("|").replace("|", " ")
    text = _BLOCKQUOTE.sub("", line)
    text = _HEADING.sub("", text)
    text = _LIST_ITEM.sub("", text)
    if block == "heading":
        text = _HEADING_TAIL.sub("", text)  # closed ATX heading: "## Title ##"
    return _EMPHASIS.sub("", text).strip()


def paragraphs_from_lines(
    lines: Sequence[str],
    *,
    page: int | None = None,
    first_line: int = 1,
    markdown: bool = False,
) -> list[Paragraph]:
    """Group lines into paragraphs, keeping each line's number and offset.

    ``first_line`` is the document line number of ``lines[0]`` and ``page`` the page
    those lines came from, so a PDF page can be ingested with the same code.
    """
    blocks: list[list[tuple[int, str]]] = []
    current: list[tuple[int, str]] = []

    def flush() -> None:
        nonlocal current
        if current:
            blocks.append(current)
            current = []

    def consume(number: int, raw: str) -> None:
        """Take one line that is known not to open or close a code fence."""
        if not raw.strip() or (markdown and _RULE.match(raw)):
            flush()  # a blank line or a horizontal rule separates, it never speaks
            return
        block: Block = _block_of(raw) if markdown else "text"
        if block != "text":
            flush()
        text = _clean(raw, block=block, markdown=markdown)
        if block == "row" and not any(character.isalnum() for character in text):
            return  # a table's "|---|---|" separator row is not prose
        if not text:
            return
        # A bibliography heading stands alone in every mode, blank line or not:
        # papers run the entries straight under it, and `find_bibliography` has to
        # see the heading by itself.
        heading = bool(_BIBLIOGRAPHY_HEADING.match(text))
        if heading:
            flush()
        current.append((number, text))
        if heading or block in _STANDALONE:
            flush()

    fenced = False
    held: list[tuple[int, str]] = []  # lines inside a code fence that is still open
    for position, raw in enumerate(lines):
        number = first_line + position
        if fenced:
            if _FENCE.match(raw):
                fenced, held = False, []  # closed: it was code, and code is skipped whole
            else:
                held.append((number, raw))
            continue
        if markdown and _FENCE.match(raw):
            flush()
            fenced, held = True, []
            continue
        consume(number, raw)
    flush()
    # A fence that never closes was no code block, and its lines may not be dropped
    # (they are the tail of the document, so replaying them keeps the file's order).
    # The fence line itself stays out: it is markup, not content.
    for number, raw in held:
        consume(number, raw)
    flush()
    return [_paragraph(index, block, page) for index, block in enumerate(blocks)]


def _append(text: str, addition: str) -> tuple[str, int]:
    """Append one line to a paragraph, returning the text and where the line starts."""
    if not text:
        return addition, 0  # the first line always starts at offset 0
    if _WORD_HYPHEN.search(text):
        # A word broken across lines is one word: "inter-" + "national". The
        # hyphen goes only before a lower-case continuation; "German-" + "American"
        # keeps it, because there it is the word's own hyphen. A hyphen standing on
        # its own ("the trial -") joins neither way and keeps its spaces.
        if addition[:1].islower():
            text = text[:-1]
        return text + addition, len(text)
    text += " "
    return text + addition, len(text)


def _paragraph(index: int, block: Sequence[tuple[int, str]], page: int | None) -> Paragraph:
    """Join one block's lines into a paragraph, recording where each line lands."""
    text = ""
    positions: list[tuple[int, int]] = []
    for number, line in block:
        text, offset = _append(text, line)
        positions.append((number, offset))
    return _build(index, text, positions, page)


def _build(index: int, text: str, lines: Sequence[tuple[int, int]], page: int | None) -> Paragraph:
    """Wrap a flattened paragraph and its line map in the value type."""
    paragraph_lines = tuple(lines)
    sentences = tuple(
        Sentence(text=text[start:end], locator=_locate(paragraph_lines, start, page),
                 start=start, end=end)
        for start, end in sentence_spans(text)
    )  # fmt: skip
    return Paragraph(
        index=index,
        text=text,
        locator=Locator(line=paragraph_lines[0][0], column=1, page=page),
        sentences=sentences,
        lines=paragraph_lines,
    )


def _locate(lines: Sequence[tuple[int, int]], offset: int, page: int | None) -> Locator:
    """Map an offset in a paragraph's text to a line. Same rule as ``Document.locate``."""
    position = bisect_right([start for _, start in lines], offset) - 1
    line, line_start = lines[position]
    return Locator(line=line, column=offset - line_start + 1, page=page)


def find_bibliography(paragraphs: Sequence[Paragraph]) -> int | None:
    """Index of the heading paragraph the bibliography starts under, or ``None``.

    A running head can repeat "References" on every page, so the last heading that
    actually has something under it wins.
    """
    for index in range(len(paragraphs) - 2, -1, -1):  # -2: entries must follow it
        if _BIBLIOGRAPHY_HEADING.match(paragraphs[index].text.strip()):
            return index
    return None


# "[12] ", "12. ", "12) ", or a printed number set off by a wide gap. A bare number is
# capped at three digits: "2019. BERT: ..." opens an author-year entry with a year,
# and reading that as entry number 2019 loses the whole bibliography. A bracketed
# number has no such ambiguity, so it keeps its four digits.
_ENTRY_START = re.compile(r"^\s*(?:\[(\d{1,4})\]|(\d{1,3})[.)]\s|(\d{1,3})\s{2,})")
# Nothing constrains where a bare list may *open*: a references section printed from
# page 4 of a split file starts at 11, and renumbering it by position would point "[12]"
# at the entry printed 11 -- the wrong source, checked in silence. What a bare number has
# to do is keep counting upward (see `accepted` below); the three-digit cap above is what
# keeps a year out.


# What a run of numbered paragraphs has to reach before it may be read as a
# bibliography nobody put a heading on. Four numbered lines are a numbered list; a
# reference list is longer than that, and mistaking one for the other would point
# every citation at the wrong source.
MIN_FALLBACK_ENTRIES = 5

# offset in the paragraph text, line number, printed number, marker style.
Cut = tuple[int, int, int, str]


def _entry_cuts(paragraph: Paragraph) -> list[Cut]:
    """Every line of ``paragraph`` that could start a numbered entry."""
    offsets = [offset for _, offset in paragraph.lines]
    cuts: list[Cut] = []
    for position, (number, offset) in enumerate(paragraph.lines):
        end = offsets[position + 1] if position + 1 < len(offsets) else len(paragraph.text)
        match = _ENTRY_START.match(paragraph.text[offset:end])
        if match is None:
            continue
        printed = next((int(group) for group in match.groups() if group is not None), 0)
        # "[12]" and "12." are different conventions; a list uses one of them.
        style = "bracket" if match.group(1) else "bare"
        cuts.append((offset, number, printed, style))
    return cuts


def find_last_numbered_run(paragraphs: Sequence[Paragraph]) -> tuple[int, int] | None:
    """Half-open range of the last run of numbered entries, or ``None``.

    Nature and its family print the reference list with no heading of any kind, so
    ``find_bibliography`` finds nothing and the paper arrives with no references at
    all -- its citation markers then have nothing to be checked against, which the
    coverage block can only report as a gap (OPEN-ITEMS 9.3 / 11.4).

    What is left to go on is the shape of the block: a contiguous run of paragraphs
    whose lines open with a printed number, the numbers counting upward, and enough
    of them that a numbered list in the prose cannot be mistaken for one. The *last*
    such run, because a paper's methods may number its steps and its bibliography
    still comes after them.
    """
    marked = [index for index, paragraph in enumerate(paragraphs) if _entry_cuts(paragraph)]
    if not marked:
        return None
    numbered = set(marked)
    start = marked[-1]
    while start - 1 in numbered:
        start -= 1
    printed = [
        number
        for index in range(start, marked[-1] + 1)
        for _, _, number, _ in _entry_cuts(paragraphs[index])
    ]
    if len(printed) < MIN_FALLBACK_ENTRIES:
        return None
    # Strictly ascending: a bibliography counts up, a protocol's steps or a table of
    # measurements do not have to. A gap is fine -- OCR loses entries -- but a number
    # that goes backwards means these are not entries in one list.
    if any(later <= earlier for earlier, later in pairwise(printed)):
        return None
    # The range ends at the last numbered paragraph, not at the end of the document.
    # A heading says "everything after me is the bibliography"; a run found by shape
    # says only what it covers, and whatever a journal prints after the list --
    # acknowledgements, a reporting summary, author contributions -- is body text
    # that would otherwise be glued onto the final entry.
    return (start, marked[-1] + 1)


def _collapse(text: str) -> str:
    return " ".join(text.split())


def split_references(paragraphs: Sequence[Paragraph], *, start: int) -> list[Reference]:
    """Split the bibliography into entries. ``start`` is the first entry paragraph.

    Entries are cut on lines that carry a printed number; a line without one
    continues the entry above it. A bibliography with no numbering at all falls
    back to one entry per paragraph, which is how most author-year lists are set.
    """
    block = list(paragraphs[start:])
    if not block:
        return []
    # No candidate at all: the list carries no usable numbers and position is all a
    # citation could mean. One entry per paragraph, numbered by position.
    if not any(_entry_cuts(paragraph) for paragraph in block):
        return [
            Reference(number=ordinal + 1, raw=_collapse(paragraph.text), locator=paragraph.locator)
            for ordinal, paragraph in enumerate(block)
        ]

    style: str | None = None
    previous = 0

    def accepted(cut: Cut) -> bool:
        """Does this candidate really start the next entry, or continue the last one?"""
        nonlocal style, previous
        _, _, printed, cut_style = cut
        if style is None:  # the first marker sets the list's convention
            style, previous = cut_style, printed
            return True
        if cut_style != style:
            return False
        # "[5]" can be nothing but a marker, so a gap in the numbering is none of our
        # business. A bare "2020." is a year at least as often as a marker, so it has
        # to keep counting upward to be believed -- a short gap is an OCR loss, a jump
        # is a date.
        if style == "bare" and not previous < printed <= previous + 9:
            return False
        previous = printed
        return True

    entries: list[tuple[int, Locator, list[str]]] = []
    for paragraph in block:
        page = paragraph.locator.page
        cuts = [cut for cut in _entry_cuts(paragraph) if accepted(cut)]
        bounds = [offset for offset, _, _, _ in cuts] + [len(paragraph.text)]
        head = paragraph.text[: bounds[0]].strip()
        # Text before the first number continues the entry above it. With no entry
        # above, it is heading residue ("Notes", a running head) and never an entry:
        # numbering it would put a second reference 1 next to the printed "[1]".
        if head and entries:
            entries[-1][2].append(head)
        for position, (offset, number, printed, _) in enumerate(cuts):
            # One paragraph, one page. An entry rejoined across a page break keeps the
            # page it started on, which is where a reader looks for its number; an
            # entry that starts inside the rejoined half would inherit that page too,
            # which is accepted -- a Locator names one page (see `_merge_across_pages`).
            piece = paragraph.text[offset : bounds[position + 1]].strip()
            entries.append((printed, Locator(line=number, column=1, page=page), [piece]))
    return [
        Reference(
            # The printed number is what "[12]" in the body points at; without one
            # the position in the list is all a citation can mean.
            number=printed if printed >= 1 else ordinal + 1,
            raw=_collapse(" ".join(pieces)),
            locator=locator,
        )
        for ordinal, (printed, locator, pieces) in enumerate(entries)
    ]


def _assemble(
    paragraphs: Sequence[Paragraph],
    *,
    name: str,
    kind: Kind,
    pages: int,
    errors: Sequence[PageError] = (),
    across_pages: bool = False,
    headless_fallback: bool = False,
) -> Document:
    """Cut the bibliography off the body and put the Document together.

    ``across_pages`` rejoins an entry whose printing ran over a page break; only a
    paged format asks for it. ``headless_fallback`` lets a paged document whose
    reference list carries no heading be found by its shape instead (see
    ``find_last_numbered_run``); a pasted draft keeps its numbered lists.
    """
    heading = find_bibliography(paragraphs)
    span: tuple[int, int] | None = None
    if heading is not None:
        span = (heading + 1, len(paragraphs))
    elif headless_fallback:
        span = find_last_numbered_run(paragraphs)
    if span is None:
        body, references = list(paragraphs), []
    else:
        start, end = span
        # The heading itself is dropped; a run found by shape has none to drop.
        body = list(paragraphs[: heading if heading is not None else start])
        body.extend(paragraphs[end:])
        entries = list(paragraphs[start:end])
        references = split_references(
            _merge_across_pages(entries) if across_pages else entries, start=0
        )
    return Document(
        name=name,
        kind=kind,
        # The body is re-indexed: a Paragraph's index is its position in the tuple.
        paragraphs=tuple(replace(p, index=position) for position, p in enumerate(body)),
        references=tuple(references),
        pages=pages,
        errors=tuple(errors),
    )


def _merge_across_pages(paragraphs: Sequence[Paragraph]) -> list[Paragraph]:
    """Rejoin a bibliography entry that the printer split over a page break.

    Only the bibliography is put back together. A body paragraph broken by a page
    break stays two paragraphs: its second half sits on its own page and carries its
    own line numbers, and a Locator names one page.
    """
    merged: list[Paragraph] = []
    previous_page: int | None = None
    for paragraph in paragraphs:
        page = paragraph.locator.page
        continues = (
            bool(merged)
            and page != previous_page  # only the first paragraph of a page continues one
            and not merged[-1].text.rstrip().endswith(".")  # a closed entry continues nothing
            and _ENTRY_START.match(paragraph.text) is None  # a marker starts a new entry
        )
        if continues:
            merged[-1] = _join(merged[-1], paragraph)
        else:
            merged.append(paragraph)
        previous_page = page
    return merged


def _join(first: Paragraph, second: Paragraph) -> Paragraph:
    """One paragraph printed in two pieces, put back together under the first's page."""
    text, offset = _append(first.text, second.text)
    lines = list(first.lines) + [(number, start + offset) for number, start in second.lines]
    return _build(first.index, text, lines, first.locator.page)


def from_text(text: str, *, name: str, kind: Kind, markdown: bool = False) -> Document:
    """Build a Document from one string: body paragraphs first, bibliography after."""
    paragraphs = paragraphs_from_lines(text.splitlines(), markdown=markdown)
    return _assemble(paragraphs, name=name, kind=kind, pages=1)


#: A link a post carries that is the post's own address, and which is therefore left
#: out of the bibliography. The reader is told rather than left to notice: one of the
#: post's own links is missing from the report, and this says which and why.
SELF_LINK_DROPPED = "a link back to the post itself was dropped: {url}"


def _address_key(url: str) -> str:
    """What a post address *names*, for telling one post's link from its own address.

    Only the parts that choose a post are kept. Host case and a ``www.`` prefix say
    nothing about which post is meant -- every matcher in ``providers.social`` reads a
    host that way -- and neither does the scheme, a trailing slash, or a fragment
    naming a place inside the same post. The path keeps its own case: a Bluesky rkey
    and a Mastodon status id are case-sensitive, and two rkeys that differ only in
    case are two posts. The query stays because it is where a Hacker News item's
    number is written. An address with no host at all matches nothing.
    """
    split = urlsplit(url.strip())
    host = (split.hostname or "").lower().removeprefix("www.")
    if not host:
        return ""
    try:
        port = f":{split.port}" if split.port else ""
    except ValueError:
        # A netloc whose port is not a number names no host anyone can reach; the
        # rest of the address is still compared, and the malformed port is ignored.
        port = ""
    query = f"?{split.query}" if split.query else ""
    return f"{host}{port}{split.path.rstrip('/')}{query}"


def from_post(post: SocialPost) -> tuple[Document, tuple[str, ...]]:
    """A post and the posts it quotes, as a document whose references are its links.

    Spec section 6.2: a post cites by linking, so the addresses inside it are the
    bibliography and the post's own words are what stands on them. One paragraph per
    post, the quoted ones after the quoter in the order they were read, and one
    :class:`~proofpath.document.Reference` per link numbered in order of appearance
    across the whole thing.

    Each link's locator is the first line of the post that carried it, which is what
    ``claims.pair_links`` pairs on: a quoted post's sources are its own author's, and
    must not end up backing the words of whoever quoted it.

    A link that points back at the post carrying it is not a source and is dropped.
    Kept, it would be resolved like any other reference, routed back to
    ``SocialProvider.fetch``, and the post would be read a second time -- handing its
    own sentence back as the passage that supports it, which is product rule 1's
    whole subject. Posts reach this shape by ordinary means: a Bluesky facet or link
    card pointing at its own permalink, a Mastodon anchor, and a Hacker News story
    whose submitted ``url`` is the item itself. ``providers.social`` already refuses
    the same thing for a Reddit self post, whose ``url`` is the post; this is that
    rule where every platform passes. The returned notes name each dropped link, so a
    reader comparing the report against the post can see why one is not in it.
    """
    lines: list[str] = []
    links: list[tuple[str, int]] = []
    notes: list[str] = []
    for item in (post, *post.quoted):
        # A post with no words gets no paragraph -- there would be no sentence for a
        # claim to be made of -- but its links are still listed, so nothing it cited
        # disappears between the reader and the report (product rule 6). They stay
        # uncited: there is no sentence of this author's for them to stand behind.
        body = [line for line in item.text.splitlines() if line.strip()]
        if body:
            if lines:
                lines.append("")  # one blank line is what separates two paragraphs
            first = len(lines) + 1
            lines.extend(body)
        else:
            # A blank line, which no paragraph owns. Recording the line the *next*
            # post's body will start on instead would hand these links to that post's
            # paragraph, and ``claims.pair_links`` would then back the quoted author's
            # sentences with a source only the quoter posted (product rule 1). A
            # Bluesky ``recordWithMedia#view`` is exactly that shape: a quote and a
            # link card, with the quoter writing nothing of their own.
            lines.append("")
            first = len(lines)
        # Per post, not per document: the quoter's link to the post it quotes is a
        # link to another author's words, and only a post's link to *itself* is the
        # one that would come back as its own evidence.
        own = _address_key(item.url)
        for link in item.links:
            if own and _address_key(link) == own:
                notes.append(SELF_LINK_DROPPED.format(url=link))
                continue
            links.append((link, first))
    document = Document(
        name=post.url,
        kind="post",
        paragraphs=tuple(paragraphs_from_lines(lines)),
        references=tuple(
            Reference(number=number, raw=url, locator=Locator(line=line))
            for number, (url, line) in enumerate(links, start=1)
        ),
        pages=1,
    )
    return document, tuple(notes)


def with_link_references(document: Document) -> Document:
    """The same document, with the addresses in its body as its bibliography.

    For pasted text that prints no reference list: the sources of a paragraph a
    reader pasted out of a post are the links inside it (spec section 6.2). One
    reference per distinct address, numbered in order of first appearance, each
    located where it was first written -- which is the paragraph whose sentences it
    then backs.
    """
    found: dict[str, Locator] = {}
    for paragraph in document.paragraphs:
        for url, offset in resolve.find_urls(paragraph.text):
            if url not in found:
                found[url] = document.locate(paragraph.index, offset)
    if not found:
        return document
    return replace(
        document,
        references=tuple(
            Reference(number=number, raw=url, locator=locator)
            for number, (url, locator) in enumerate(found.items(), start=1)
        ),
        # Said in the kind, and only here: the caller asked for the addresses to *be*
        # the bibliography, which is the one thing that makes ``claims.pair_links``
        # right about this text. A document that merely happens to print an address
        # is not this, and must never be paired that way (product rule 1).
        kind="linked",
    )


def _read(path: Path) -> str:
    try:
        # errors="replace": a mis-encoded byte costs one character, not the document.
        return path.read_text(encoding="utf-8", errors="replace")
    except OSError as error:
        raise IngestError(f"cannot read {path}: {error}") from error


def from_markdown(path: Path) -> Document:
    """Ingest a markdown file."""
    return from_text(_read(path), name=path.name, kind="markdown", markdown=True)


def from_plain(path: Path) -> Document:
    """Ingest a plain text file."""
    return from_text(_read(path), name=path.name, kind="text")


# A block that is nothing but a page number is furniture -- but only in the margin:
# "42" halfway down the page is a measurement, and dropping it would lose a claim.
_PAGE_NUMBER = re.compile(r"^\d+$")
_MARGIN = 0.1  # top and bottom tenth of the page
# A short block printed at the same height on this many pages is a running head.
_RUNNING_PAGES = 3
_RUNNING_WORDS = 8


@dataclass(frozen=True)
class _Block:
    """One text block of a PDF page: where it sits, and the lines it holds."""

    page: int
    top: int  # round(y0): a running head sits at the same height on every page
    left: float
    lines: tuple[str, ...]
    first_line: int = 1  # the number of ``lines[0]`` on its page
    margin: bool = False  # in the top or bottom tenth of the page, where furniture sits

    @property
    def text(self) -> str:
        """The block as one line, for recognising a repeat across pages."""
        return _collapse(" ".join(self.lines))


# pymupdf's span flag bit 0: this span is printed as a superscript.
_SUPERSCRIPT = 1
# A superscript holding nothing but a citation: "51", "3,4", "4-6" and their dashes.
_SUPERSCRIPT_CITATION = re.compile(r"^\d{1,3}(?:\s*[,;\u2013\u2014-]\s*\d{1,3})*$")
# What the superscript sits on decides what it is. Only a standalone number takes an
# exponent ("10" -> ten cubed); digits that end a name do not ("Uniclust30", "PDB70",
# "6Y4F" are all cited in the AlphaFold paper), and neither does a space in front of the
# superscript. Anything else is read as a citation, because losing a real one is the
# failure this rewrite exists to prevent. A full stop in the lookbehind is what keeps
# "OpenMM v.7.3.1" a version rather than a base: the "1" ends a dotted field, not a
# number. Its cost is that "0.5" followed by a superscript is read as a citation too --
# a false marker a reader can see, traded against a real citation nobody would.
_EXPONENT_BASE = re.compile(r"(?<![A-Za-z0-9.])\d+$")
# A superscript printed *after* punctuation is a footnote marker, not a citation: ACL and
# Chicago set footnotes that way ("...and code.1", "...training,8"), while Vancouver and
# Nature bind a citation to the word it belongs to ("...prediction51."). Reading a
# footnote marker as a citation would assert that reference 1 supports a sentence that
# never named it -- a worse failure than the invisibility this rewrite exists to fix. A
# closing bracket is not punctuation for this purpose: "(PDB)5" is how Nature prints a
# citation on a parenthesised term. Measured on the three papers in
# docs/eval/2026-09-11-pairing.md: counted by hand against the printed pages, this drops
# 10 footnote markers out of 10 and keeps 116 citations out of 116.
_FOOTNOTE_BASE = re.compile(r"[.,;:!?]$")
# Whitespace and closing quotes sit between the stop and the footnote marker without
# changing what it is; a closing bracket is not in here on purpose.
_CLOSERS = " \t\u201d\u2019\"'"


def _line_text(spans: Sequence[dict[str, Any]]) -> str:
    """One PDF line's spans joined, with superscript citation markers spelled out.

    Nature-style citations are printed as superscripts, and flattening a page to text
    glues them onto the word before ("...structure prediction51"), where nothing can
    see them: the paper then looks like a paper with no citations, which is exactly the
    silence spec section 15 forbids. Where the PDF itself marks a span as superscript
    and the span holds nothing but numbers, it is written out as "[51]" -- the same
    marker the same citation would carry in a bracketed journal. A superscript the PDF
    does not mark is still invisible; nothing here guesses from position on the page.
    """
    text = ""
    for span in spans:
        piece = str(span.get("text", ""))
        stripped = piece.strip()
        # Two different views of what comes before. An exponent is printed tight against
        # its base, so the raw text decides it -- a space in front means citation. A
        # footnote marker may sit behind a closing quote ("...herself." 1), so spaces and
        # quotes are stripped before the punctuation is read.
        if (
            int(span.get("flags", 0)) & _SUPERSCRIPT
            and _SUPERSCRIPT_CITATION.fullmatch(stripped)
            and not _EXPONENT_BASE.search(text)
            and not _FOOTNOTE_BASE.search(text.rstrip(_CLOSERS))
        ):
            piece = f"[{stripped}]"
        text += piece
    return text


_TEXT_BLOCK = 0
_IMAGE_BLOCK = 1
# A page of images and no text is a scan: it holds claims a reader can see and this
# version cannot read. Saying so is the whole of product rule 2 -- an unread page must
# never arrive as an empty one. A page with neither (a separator) hides nothing.
SCANNED_PAGE = "no extractable text; the page is an image (scanned?)"


def _page_blocks(page: pymupdf.Page, number: int) -> tuple[list[_Block], bool]:
    """Every text block of one page in reading order, and whether the page is an image.

    The flag is true only when the page carries at least one image block and no text
    block at all: that is the shape of a scan, and it is reported as a ``PageError``.
    """
    blocks: list[_Block] = []
    # pymupdf ships py.typed but leaves ``get_text`` unannotated.
    blocks_of_page = page.get_text("dict")["blocks"]  # type: ignore[no-untyped-call]
    types = {block["type"] for block in blocks_of_page}
    image_only = _IMAGE_BLOCK in types and _TEXT_BLOCK not in types
    for block in blocks_of_page:
        if block["type"] != _TEXT_BLOCK:
            continue  # an image cannot carry a claim; only text can
        lines = tuple(
            text for line in block["lines"] if (text := _line_text(line["spans"]).strip())
        )
        if not lines:
            continue
        x0, y0, _, y1 = block["bbox"]
        height = page.rect.height
        blocks.append(
            _Block(
                page=number,
                top=round(y0),
                left=x0,
                lines=lines,
                margin=y0 < _MARGIN * height or y1 > (1 - _MARGIN) * height,
            )
        )
    blocks.sort(key=lambda block: (block.top, block.left))
    # Number first, filter second: "p.4 L3" is the third text line a reader counts on
    # page 4, and a running head that is dropped further down was still one of them.
    numbered: list[_Block] = []
    line = 1  # a PDF locator counts lines within the page, not within the file
    for block in blocks:
        numbered.append(replace(block, first_line=line))
        line += len(block.lines)
    return numbered, image_only


def _without_furniture(pages: Sequence[Sequence[_Block]]) -> list[list[_Block]]:
    """Drop page numbers and running heads. Collect first, filter second.

    A repeat cannot be recognised from one page, so every page is read before
    anything is dropped.
    """
    seen: dict[tuple[int, str], set[int]] = {}
    for blocks in pages:
        for block in blocks:
            seen.setdefault((block.top, block.text), set()).add(block.page)

    def furniture(block: _Block) -> bool:
        text = block.text
        if block.margin and _PAGE_NUMBER.match(text):
            return True
        # Length is the guard: a paragraph that happens to repeat is still prose.
        return len(text.split()) < _RUNNING_WORDS and len(seen[(block.top, text)]) >= _RUNNING_PAGES

    return [[block for block in blocks if not furniture(block)] for blocks in pages]


def _paragraphs_of_pages(pages: Sequence[Sequence[_Block]]) -> list[Paragraph]:
    """Paragraphs in page order, each block keeping the line numbers it was given."""
    return [
        paragraph
        for blocks in pages
        for block in blocks
        for paragraph in paragraphs_from_lines(
            block.lines, page=block.page, first_line=block.first_line
        )
    ]


def from_pdf(path: Path) -> Document:
    """Ingest a PDF: paragraphs keep their page, and a page that fails is reported.

    A superscript citation marker is recognised where the PDF marks the span as one
    (``_line_text``), and is written out as "[12]" so it can be paired like any other;
    a marker the PDF leaves unmarked still has to be printed as text to be seen.

    A page that holds an image and no text is a scan and is reported as a ``PageError``
    rather than passed on as a page with nothing on it.
    """
    try:
        # ``pymupdf.open`` is the ``Document`` constructor, which ships unannotated.
        opened = pymupdf.open(path)  # type: ignore[no-untyped-call]
    except (RuntimeError, ValueError, OSError) as error:
        # pymupdf's FileDataError, FileNotFoundError and EmptyFileError are all
        # RuntimeError subclasses; a document nobody can open is not a page error.
        raise IngestError(f"cannot open {path}: {error}") from error

    errors: list[PageError] = []
    pages: list[list[_Block]] = []
    with opened as document:
        count = document.page_count
        if count < 1:
            raise IngestError(f"cannot read {path}: the document has no pages")
        for index in range(count):
            number = index + 1
            try:
                blocks, image_only = _page_blocks(document[index], number)
            except (RuntimeError, ValueError) as error:
                # Spec section 15: the page is named and the run goes on. An
                # unreadable page is never reported as an empty one.
                errors.append(PageError(page=number, detail=f"{type(error).__name__}: {error}"))
                continue
            pages.append(blocks)
            if image_only:
                errors.append(PageError(page=number, detail=SCANNED_PAGE))
    return _assemble(
        _paragraphs_of_pages(_without_furniture(pages)),
        name=path.name,
        kind="pdf",
        pages=count,
        errors=errors,
        across_pages=True,
        headless_fallback=True,
    )


def _docx_lines(document: docx.document.Document) -> list[tuple[int, str]]:
    """The body in document order: ``(ordinal, text)`` for everything that has text.

    The ordinal counts every body-level paragraph, empty ones included, and every
    table row, so the numbering stays stable however much of the file is blank.
    """
    lines: list[tuple[int, str]] = []
    ordinal = 0
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            for row in item.rows:
                ordinal += 1
                # A row is one line: the cells are read across, as they are printed.
                cells = [text for cell in row.cells if (text := cell.text.strip())]
                if cells:
                    lines.append((ordinal, " | ".join(cells)))
            continue
        ordinal += 1
        if item.text.strip():
            lines.append((ordinal, item.text))
    return lines


def from_docx(path: Path) -> Document:
    """Ingest a .docx. Lines are numbered by body paragraph and table row.

    Each Word paragraph is a paragraph here: Word puts no blank paragraph between
    them -- the spacing is a style -- so merging consecutive ones would flatten a page
    of prose into one claim. Headers and footers are not read at all: they are
    furniture by definition, the same running heads `from_pdf` drops.

    A .docx has no fixed pages -- pagination is the renderer's -- so locators carry a
    line and no page.
    """
    try:
        opened = docx.Document(str(path))
    # A .docx is a zip of XML: a truncated one, a zip that is not an OPC package, and
    # a package missing a part all show up differently, and none of them is a document.
    except (OpcError, zipfile.BadZipFile, KeyError, OSError, ValueError) as error:
        raise IngestError(f"cannot open {path}: {error}") from error
    paragraphs = [
        paragraph
        for ordinal, text in _docx_lines(opened)
        for paragraph in paragraphs_from_lines([text], first_line=ordinal)
    ]
    return _assemble(paragraphs, name=path.name, kind="docx", pages=1)


_LOADERS: dict[str, Callable[[Path], Document]] = {
    ".pdf": from_pdf,
    ".docx": from_docx,
    ".md": from_markdown,
    ".markdown": from_markdown,
    ".txt": from_plain,
}


def load(target: Path) -> Document:
    """Ingest a file, choosing the reader by suffix. The entry point for the CLI."""
    suffix = target.suffix.lower()
    loader = _LOADERS.get(suffix)
    if loader is None:
        raise IngestError(
            f"unsupported file type: {suffix or target.name} (expected .pdf, .docx, .md, .txt)"
        )
    if not target.is_file():
        raise IngestError(f"cannot read {target}: no such file")
    return loader(target)
